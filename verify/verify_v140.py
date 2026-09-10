"""v140 verification: audit fixes (atomic writes, kit validation, dup function),
simple house-style cover titles, configurable brand, and the Excel review sheet."""
import sys, shutil, tempfile, inspect
from pathlib import Path

REPO, BK = Path(sys.argv[1]), Path(sys.argv[2])
sys.path.insert(0, str(REPO))
from PIL import Image
from pypdf import PdfReader
from reportlab.pdfgen import canvas as rlc
from reportlab.pdfbase import pdfmetrics

from docrefine import reviews, classify
from docrefine.rebrand import (BrandKit, rebrand_pdf, title_for, title_from_filename,
                               fit_title, _ensure_font, _FONT_NAME, TITLE_FRAC,
                               ASSET_TYPE_TITLES, DEFAULT_TITLE)
from docrefine.worker import Worker

results = []
def check(n, ok, d=""):
    results.append(bool(ok)); print(f"[{'PASS' if ok else 'FAIL'}] {n}  {d}")

def text_pdf(path, text="SAMPLE BODY TEXT FOR TESTING"):
    path.parent.mkdir(parents=True, exist_ok=True)
    c = rlc.Canvas(str(path), pagesize=(612, 792)); c.setFont("Helvetica", 22)
    c.drawString(70, 700, text); c.showPage(); c.save()

def W(): return Worker(callback=lambda e: None)

work = Path(tempfile.mkdtemp(prefix="drp_v140_"))
sheets = []
kit = BrandKit(BK)
_ensure_font()

# =====================================================================
#  C1 — an interrupted write must never leave a usable-looking file
# =====================================================================
d = work / "atomic"; d.mkdir(parents=True)
text_pdf(d / "in.pdf")
import docrefine.rebrand as R
orig_write = R.PdfWriter.write
def crash(self, f):
    f.write(b"%PDF-1.4 partially written"); raise IOError("simulated crash")
R.PdfWriter.write = crash
try:
    rebrand_pdf(d / "in.pdf", d / "out.pdf", kit, "Installation Manual")
except Exception:
    pass
R.PdfWriter.write = orig_write
check("C1 crash mid-write leaves no output file", not (d / "out.pdf").exists())
check("C1 no .part debris left behind", not list(d.glob("*.part")))

rebrand_pdf(d / "in.pdf", d / "out.pdf", kit, "Installation Manual")
check("C1 a normal write still produces a valid PDF",
      (d / "out.pdf").exists() and len(PdfReader(str(d / "out.pdf")).pages) == 3)

# _atomic_copy keeps the same guarantee for copied files
from docrefine.worker import _atomic_copy
src_f = d / "src.bin"; src_f.write_bytes(b"x" * 5000)
_atomic_copy(src_f, d / "copy.bin")
check("C1 atomic copy is byte-identical", (d / "copy.bin").read_bytes() == src_f.read_bytes())

# =====================================================================
#  C2 — one ollama_available, and it is the auto-starting one
# =====================================================================
srctext = (REPO / "docrefine" / "classify.py").read_text(encoding="utf-8")
check("C2 ollama_available defined exactly once", srctext.count("def ollama_available") == 1,
      f"{srctext.count('def ollama_available')} definitions")
check("C2 the surviving one auto-starts the server",
      "start_server" in inspect.getsource(classify.ollama_available))

# =====================================================================
#  C3 — a kit missing required art is rejected up front, clearly
# =====================================================================
bad = work / "badkit"; (bad / "Portrait").mkdir(parents=True)
Image.new("RGBA", (2550, 300), (10, 20, 80, 255)).save(bad / "Portrait" / "header.png")
bk = BrandKit(bad)
check("C3 kit with no cover/back is not 'has'", not bk.has("portrait"))
check("C3 missing() names the gap", set(bk.missing("portrait")) == {"cover", "back"}, str(bk.missing("portrait")))
msg = Worker._kit_problem(bk)
check("C3 error explains the real problem", "cover" in msg and "Portrait" in msg, msg[:88])
check("C3 a good kit still validates", kit.has("portrait") and kit.has("landscape"))

# =====================================================================
#  H1 — long titles wrap instead of running off the page
# =====================================================================
cov = kit.live_readers("portrait")["cover"]; cw, ch = cov._px_size
scale = min(612 / cw, 792 / ch); dh = ch * scale
nominal = TITLE_FRAC * dh; maxw = (cw * scale) * 0.76
LONG = ("INSTALLATION AND ASSEMBLY INSTRUCTIONS FOR THE VERSATILE 4C FRONT LOAD "
        "HORIZONTAL CLUSTER BOX UNIT MODEL 4C12D-22")
lines, pt = fit_title(LONG, nominal, maxw)
widest = max(pdfmetrics.stringWidth(l, _FONT_NAME, pt) for l in lines)
check("H1 long title wraps to <= 2 lines", 1 < len(lines) <= 2, f"{len(lines)} lines")
check("H1 every line fits the page", widest <= maxw, f"widest {widest:.0f} / {maxw:.0f} pt")
one, pt1 = fit_title("INSTALLATION MANUAL", nominal, maxw)
check("H1 a short title stays one line at full size", len(one) == 1 and abs(pt1 - nominal) < 0.01)
check("H1 an absurd single word is truncated, not overflowed",
      pdfmetrics.stringWidth(fit_title("X" * 400, nominal, maxw)[0][0],
                             _FONT_NAME, fit_title("X" * 400, nominal, maxw)[1]) <= maxw)

# =====================================================================
#  Titles — the Batch 1 house style: short, plain, no model numbers
# =====================================================================
check("T1 asset types map to plain titles",
      title_for("installation-guide") == "Installation Manual"
      and title_for("spec-sheet") == "Specification Sheet"
      and title_for("drawing") == "Technical Drawing"
      and title_for("warranty") == "Product Warranty")
check("T2 unknown asset type is humanised", title_for("mounting-bracket-kit") == "Mounting Bracket Kit")
check("T3 doc_type is the second chance", title_for("", doc_type="spec sheet") == "Specification Sheet")
check("T4 no-text rows get the generic title, not a mangled filename",
      classify._fallback("4C11D-09CS.pdf")["title"] == DEFAULT_TITLE)
check("T5 model numbers keep their casing",
      title_from_filename("4C11D-09CS") == "4C11D 09CS"
      and title_from_filename("tech-3635RL") == "Tech 3635RL")
check("T6 titles are short enough to read",
      all(len(v) <= 30 for v in ASSET_TYPE_TITLES.values()),
      f"longest {max(len(v) for v in ASSET_TYPE_TITLES.values())} chars")

# =====================================================================
#  Instruction flip — the 53 no-text guides that were being skipped
# =====================================================================
F = classify.filename_suggests_instructions
check("F1 instruction filenames flip",
      F("Bradford-Installation-Instructions.pdf") and F("barcelona_assembly_instructions.pdf")
      and F("206550INS-1400.pdf") and F("AF-D500-Install.pdf"))
check("F2 drawing filenames do not flip",
      not F("tech-3635RL.pdf") and not F("4c12d-22-cut-sheet.pdf") and not F("120RCS_BM.pdf"))
nt = work / "notext"; nt.mkdir()
blank = nt / "Widget-Installation-Instructions.pdf"
c = rlc.Canvas(str(blank), pagesize=(612, 792)); c.rect(80, 80, 300, 300, fill=1); c.showPage(); c.save()
info = classify.classify_document(blank)
check("F3 a no-text instruction PDF is flagged for rebrand", info["action"] == "rebrand", str(info["action"]))
check("F4 and carries a sensible title", info["title"] == "Installation Manual", info["title"])
blank2 = nt / "tech-9999XL.pdf"
c = rlc.Canvas(str(blank2), pagesize=(612, 792)); c.rect(80, 80, 300, 300, fill=1); c.showPage(); c.save()
check("F5 a no-text drawing still defaults to leave",
      classify.classify_document(blank2)["action"] == "leave")

# =====================================================================
#  M1 — brand really is swappable now
# =====================================================================
(bad / "brand.json").write_text('{"name": "Acme Mail Co", "slug": "Acme Mail!"}', encoding="utf-8")
ak = BrandKit(bad)
check("M1 brand.json sets the name", ak.brand_name == "Acme Mail Co", ak.brand_name)
check("M1 slug is normalised", ak.brand_slug == "acme-mail", ak.brand_slug)
check("M1 subtitle uses the kit's brand",
      ak.subtitle_for("Florence") == "Manufactured by Florence | Sold by Acme Mail Co")
check("M1 default stays Budget Mailboxes", kit.brand_name == "Budget Mailboxes"
      and kit.subtitle_for("") == "Sold by Budget Mailboxes")

# =====================================================================
#  M3 — text past page 2 is still found
# =====================================================================
deep = work / "deep.pdf"
c = rlc.Canvas(str(deep), pagesize=(612, 792))
for _ in range(2):
    c.rect(80, 80, 400, 400, fill=1); c.showPage()      # two image-only pages
c.setFont("Helvetica", 20); c.drawString(70, 700, "REAL INSTALLATION TEXT ON PAGE THREE OF THIS DOCUMENT")
c.showPage(); c.save()
check("M3 text on page 3 is picked up", len(classify.extract_text(deep)) > 20,
      f"{len(classify.extract_text(deep))} chars")

# =====================================================================
#  U4 — the Excel review sheet
# =====================================================================
xsrc = work / "xl"; text_pdf(xsrc / "a.pdf"); text_pdf(xsrc / "sub" / "b.pdf")
rows = [
    {"file": "a.pdf", "action": "rebrand", "product": "A", "asset_type": "spec-sheet",
     "confidence": "0.95", "source": "llm", "title": "Specification Sheet"},
    {"file": "sub/b.pdf", "action": "leave", "product": "B", "asset_type": "document",
     "confidence": "0", "source": "fallback", "title": DEFAULT_TITLE},
]
xp = work / "plan.xlsx"
reviews.write_plan(xp, rows, Worker.REBRAND_PLAN_COLUMNS, src_root=xsrc,
                   asset_types=sorted(ASSET_TYPE_TITLES))
check("U4 xlsx written", xp.exists() and xp.stat().st_size > 4000, f"{xp.stat().st_size} bytes")
back = reviews.read_plan(xp)
check("U4 round-trips every row", len(back) == 2, f"{len(back)} rows")
check("U4 round-trips the values",
      back[0]["file"] == "a.pdf" and back[0]["action"] == "rebrand"
      and back[1]["asset_type"] == "document", str(back[0])[:70])

from openpyxl import load_workbook
wb = load_workbook(xp); ws = wb.active
hdr = [c.value for c in ws[1]]
check("U4 triage column comes first", hdr[0] == "review?", str(hdr[:3]))
# v141 narrowed this: an unreadable file left as-is is the safe default and is
# no longer flagged. Row 3 here is fallback+leave, so neither row flags.
check("U4 confident rows are not flagged",
      not ws.cell(row=2, column=1).value and not ws.cell(row=3, column=1).value)
fcol = hdr.index("file") + 1
check("U4 filename links to the real PDF", ws.cell(row=2, column=fcol).hyperlink is not None)
check("U4 header frozen + filterable", ws.freeze_panes == "A2" and ws.auto_filter.ref is not None)
check("U4 action has a dropdown", any("action" in str(dv.sqref) or dv.formula1.find("rebrand") >= 0
                                      for dv in ws.data_validations.dataValidation))
wb.close()
check("U4 needs_review triages on the decision, not just the source (v141)",
      reviews.needs_review({"source": "fallback", "action": "rebrand", "confidence": "0"})
      and not reviews.needs_review({"source": "fallback", "action": "leave", "confidence": "0"})
      and reviews.needs_review({"source": "llm", "action": "rebrand", "confidence": "0.5"})
      and not reviews.needs_review({"source": "llm", "action": "rebrand", "confidence": "0.95"}))

# csv still readable + writable (old sheets keep working)
cp = work / "plan.csv"
reviews.write_plan(cp, rows, Worker.REBRAND_PLAN_COLUMNS)
check("U4 csv still supported both ways", len(reviews.read_plan(cp)) == 2)

# =====================================================================
#  End-to-end: analyze -> xlsx -> apply, with the brand from the kit
# =====================================================================
e2e = work / "e2e"; text_pdf(e2e / "guide.pdf", "INSTALLATION AND ASSEMBLY GUIDE FOR THE UNIT")
W().run_rebrand_analyze(str(e2e))
plan = reviews.find_plan(e2e); sheets.append(plan)
check("E1 analyze wrote an xlsx sheet", plan and plan.suffix == ".xlsx", str(plan))
prows = reviews.read_plan(plan)
for r in prows:
    r["action"] = "rebrand"; r["asset_type"] = "installation-guide"
    r["product"] = "unit"; r["title"] = ""; r["manufacturer"] = "Florence"
reviews.write_plan(plan, prows, Worker.REBRAND_PLAN_COLUMNS, src_root=e2e)
eout = work / "e2e_out"
# v143 made the attribution line opt-in; this case is about the line rendering.
W().run_rebrand_apply(str(e2e), str(BK), None, out_dir=str(eout), complete_set=True,
                      show_attribution=True)
made = list(eout.rglob("*.pdf"))
check("E2 apply consumed the xlsx sheet", len(made) == 1, str([m.name for m in made]))
if made:
    cover = (PdfReader(str(made[0])).pages[0].extract_text() or "").upper().replace(" ", "")
    check("E3 cover shows the plain house-style title", "INSTALLATIONMANUAL" in cover, cover[:60])
    check("E4 cover shows the manufacturer attribution", "FLORENCE" in cover)
    check("E5 named from the kit's brand slug",
          made[0].name.endswith("-budget-mailboxes.pdf"), made[0].name)

for s in sheets:
    try: s.unlink()
    except (OSError, AttributeError): pass
shutil.rmtree(work, ignore_errors=True)
print("\n" + "=" * 58)
print(f"RESULT: {sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
