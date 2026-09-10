import sys, glob
from openpyxl import load_workbook
from pypdf import PdfReader

REV = r"C:\Users\WORK\Documents\DocRefinePro_Data\Rebrand Reviews"

for xl in [REV + r"\Batch 4___unique-to-rebrand_rebrand_plan.xlsx",
           REV + r"\Documents__DUMP_rebrand_plan.xlsx"]:
    print("=== ", xl)
    wb = load_workbook(xl, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    print("hdr:", hdr)
    n = 0
    for r in rows:
        n += 1
        cells = [str(c) if c is not None else "" for c in r]
        if "DUMP" in xl or any("Imperial" in c for c in cells):
            print(cells)
    print("rows:", n)
    wb.close()

OUT = r"C:\Users\WORK\Documents\Batch 4\_unique-to-rebrand_rebranded"
for f in glob.glob(OUT + r"\**\*imperial-street*", recursive=True):
    try:
        rd = PdfReader(f)
        txt = "".join((p.extract_text() or "") for p in rd.pages)
        print(f"{f.split(chr(92))[-1]}: pages={len(rd.pages)} chars={len(txt)}")
    except Exception as e:
        print(f, "ERR", e)
